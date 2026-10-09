"""
SyllabiQ — Multi-Agent Academic Question-Paper Generator & Syllabus Validator
FastAPI Backend Orchestrator (Google Colab & Uvicorn Ready)

Run in Google Colab or local Python 3.11+:
  pip install fastapi uvicorn pyngrok python-multipart google-genai pymupdf python-pptx pillow
  uvicorn fastapi_main:app --host 0.0.0.0 --port 8000
"""

import os
import io
import json
import base64
from typing import List, Optional, Dict, Any
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from google import genai
from google.genai import types

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

app = FastAPI(
    title="SyllabiQ Multi-Agent FastAPI Backend",
    description="4-Agent Multimodal Question Paper Generation & Curriculum Validation API",
    version="2.4.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_client() -> genai.Client:
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY is not set.")
    return genai.Client(api_key=api_key)


def extract_pptx_slides(raw_bytes: bytes, filename: str) -> str:
    if Presentation is None:
        return f"[Presentation: {filename}]"
    prs = Presentation(io.BytesIO(raw_bytes))
    slides_out = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines = [f"--- Slide {idx} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        lines.append(txt)
        slides_out.append("\n".join(lines))
    return "\n\n".join(slides_out)


class FacultyDocumentInput(BaseModel):
    id: str
    name: str
    format: str
    mimeType: str
    base64Data: Optional[str] = None
    extractedTextPreview: Optional[str] = None


class IngestRequest(BaseModel):
    documents: List[FacultyDocumentInput]
    syllabusUnits: List[Dict[str, Any]]
    courseOutcomes: List[Dict[str, Any]]


@app.get("/api/health")
async def health():
    return {
        "status": "ok",
        "service": "SyllabiQ FastAPI Multi-Agent Engine",
        "model": "gemini-3.8-flash",
        "agents": [
            "IngestionAgent",
            "BlueprintArchitectAgent",
            "QuestionGeneratorAgent",
            "CurriculumValidatorAgent",
        ],
    }


@app.post("/api/agents/ingest")
async def run_ingestion_agent(req: IngestRequest):
    client = get_client()
    parts = []
    for doc in req.documents:
        if doc.base64Data and doc.mimeType in ["application/pdf", "image/png", "image/jpeg", "image/webp"]:
            clean_b64 = doc.base64Data.split(",")[-1]
            parts.append(types.Part.from_bytes(data=base64.b64decode(clean_b64), mime_type=doc.mimeType))
        elif doc.base64Data and doc.name.lower().endswith((".pptx", ".ppt")):
            clean_b64 = doc.base64Data.split(",")[-1]
            text = extract_pptx_slides(base64.b64decode(clean_b64), doc.name)
            parts.append(types.Part.from_text(text=f"PPTX ({doc.name}):\n{text}"))
        elif doc.extractedTextPreview:
            parts.append(types.Part.from_text(text=f"Document ({doc.name}):\n{doc.extractedTextPreview}"))

    prompt = (
        f"Analyze the uploaded faculty documents and map extracted concepts to Syllabus Units "
        f"{json.dumps(req.syllabusUnits)} and Course Outcomes {json.dumps(req.courseOutcomes)}."
    )
    parts.append(types.Part.from_text(text=prompt))
    resp = client.models.generate_content(
        model="gemini-3.8-flash",
        contents=parts,
        config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.2),
    )
    return json.loads(resp.text)
